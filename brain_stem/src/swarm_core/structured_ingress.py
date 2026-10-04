from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import math
import os
import sqlite3
import stat
from dataclasses import dataclass
from datetime import date, datetime, time
from enum import Enum
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


class StructuredIngressError(ValueError):
    pass


class StructuredFormat(str, Enum):
    CSV = "text/csv"
    JSONL = "application/x-ndjson"
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    SQLITE = "application/vnd.sqlite3"


@dataclass(frozen=True, slots=True)
class StructuredIngressLimits:
    maximum_source_bytes: int
    maximum_rows: int
    maximum_columns: int
    maximum_cell_characters: int

    def __post_init__(self) -> None:
        for field_name in (
            "maximum_source_bytes",
            "maximum_rows",
            "maximum_columns",
            "maximum_cell_characters",
        ):
            value = getattr(self, field_name)
            if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
                raise StructuredIngressError(f"{field_name} must be a positive integer")


@dataclass(frozen=True, slots=True)
class SourceCellOffset:
    source_ref: str
    source_sha256: str
    sheet_name: str | None
    row_number: int
    column_number: int
    coordinate: str


@dataclass(frozen=True, slots=True)
class ExtractedField:
    name: str
    value: object
    source_offset: SourceCellOffset


@dataclass(frozen=True, slots=True)
class ExtractedRecord:
    record_number: int
    fields: tuple[ExtractedField, ...]


@dataclass(frozen=True, slots=True)
class StructuredExtraction:
    format: StructuredFormat
    source_ref: str
    source_sha256: str
    records: tuple[ExtractedRecord, ...]


def _strict_object(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise StructuredIngressError("JSON record contains a duplicate key")
        result[key] = value
    return result


def _json_value(value: object, *, maximum_cell_characters: int) -> object:
    if value is None or isinstance(value, (str, bool, int)):
        normalized = value
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise StructuredIngressError("non-finite numeric cells are not accepted")
        normalized = value
    elif isinstance(value, (datetime, date, time)):
        normalized = value.isoformat()
    elif isinstance(value, bytes):
        normalized = {"encoding": "base64", "data": base64.b64encode(value).decode("ascii")}
    else:
        raise StructuredIngressError(f"unsupported structured cell type: {type(value).__name__}")
    try:
        encoded = json.dumps(normalized, allow_nan=False, ensure_ascii=False, separators=(",", ":"))
    except (TypeError, ValueError) as error:
        raise StructuredIngressError("cell value cannot be represented in the output contract") from error
    if len(encoded) > maximum_cell_characters:
        raise StructuredIngressError("cell value exceeds the configured character limit")
    return normalized


def _headers(values: Iterable[object], *, maximum_columns: int) -> tuple[str, ...]:
    headers = tuple(value.strip() if isinstance(value, str) else "" for value in values)
    if not headers or len(headers) > maximum_columns:
        raise StructuredIngressError("header width is empty or exceeds the configured column limit")
    if any(not header for header in headers):
        raise StructuredIngressError("headers must be non-empty strings")
    if len(set(headers)) != len(headers):
        raise StructuredIngressError("duplicate headers are not accepted")
    return headers


def _records_from_rows(
    *,
    headers: tuple[str, ...],
    rows: Iterable[tuple[int, tuple[object, ...]]],
    source_ref: str,
    source_sha256: str,
    sheet_name: str | None,
    limits: StructuredIngressLimits,
) -> tuple[ExtractedRecord, ...]:
    records: list[ExtractedRecord] = []
    for record_number, (row_number, values) in enumerate(rows, start=1):
        if record_number > limits.maximum_rows:
            raise StructuredIngressError("structured source exceeds the configured row limit")
        if len(values) != len(headers):
            raise StructuredIngressError(f"row {row_number} width does not match its header")
        fields: list[ExtractedField] = []
        for column_number, (name, raw_value) in enumerate(zip(headers, values), start=1):
            coordinate = f"{_column_name(column_number)}{row_number}"
            fields.append(
                ExtractedField(
                    name,
                    _json_value(raw_value, maximum_cell_characters=limits.maximum_cell_characters),
                    SourceCellOffset(
                        source_ref,
                        source_sha256,
                        sheet_name,
                        row_number,
                        column_number,
                        coordinate,
                    ),
                )
            )
        records.append(ExtractedRecord(record_number, tuple(fields)))
    return tuple(records)


def _column_name(column_number: int) -> str:
    name = ""
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        name = chr(65 + remainder) + name
    return name


def _read_bounded_regular_file(path: Path, maximum_bytes: int) -> bytes:
    flags = os.O_RDONLY | getattr(os, "O_CLOEXEC", 0) | getattr(os, "O_NOFOLLOW", 0)
    try:
        descriptor = os.open(path, flags)
    except OSError as error:
        raise StructuredIngressError("SQLite source contains an unreadable or symbolic-link file") from error
    try:
        file_stat = os.fstat(descriptor)
        if not stat.S_ISREG(file_stat.st_mode):
            raise StructuredIngressError("SQLite source files must be regular files")
        if file_stat.st_size > maximum_bytes:
            raise StructuredIngressError("SQLite source exceeds the configured byte limit")
        chunks: list[bytes] = []
        remaining = maximum_bytes + 1
        while remaining > 0:
            chunk = os.read(descriptor, min(1024 * 1024, remaining))
            if not chunk:
                break
            chunks.append(chunk)
            remaining -= len(chunk)
        content = b"".join(chunks)
        if len(content) > maximum_bytes:
            raise StructuredIngressError("SQLite source exceeds the configured byte limit")
        return content
    finally:
        os.close(descriptor)


class StructuredIngressService:
    def __init__(self, limits: StructuredIngressLimits) -> None:
        if not isinstance(limits, StructuredIngressLimits):
            raise TypeError("limits must be a StructuredIngressLimits")
        self._limits = limits

    def extract_bytes(
        self,
        content: bytes,
        *,
        format: StructuredFormat,
        source_ref: str,
        sheet_name: str | None = None,
    ) -> StructuredExtraction:
        if not isinstance(content, bytes):
            raise TypeError("content must be bytes")
        if len(content) > self._limits.maximum_source_bytes:
            raise StructuredIngressError("structured source exceeds the configured byte limit")
        if not isinstance(format, StructuredFormat) or format is StructuredFormat.SQLITE:
            raise StructuredIngressError("extract_bytes supports CSV, JSONL, and XLSX only")
        if not isinstance(source_ref, str) or not source_ref.strip():
            raise StructuredIngressError("source_ref must be a non-empty provenance reference")
        digest = hashlib.sha256(content).hexdigest()
        if format is StructuredFormat.CSV:
            records = self._extract_csv(content, source_ref, digest)
        elif format is StructuredFormat.JSONL:
            records = self._extract_jsonl(content, source_ref, digest)
        else:
            records = self._extract_xlsx(content, source_ref, digest, sheet_name)
        return StructuredExtraction(format, source_ref, digest, records)

    def extract_sqlite_table(
        self,
        database_path: str | Path,
        *,
        table_name: str,
        columns: tuple[str, ...],
        source_ref: str,
    ) -> StructuredExtraction:
        path = Path(database_path)
        if path.is_symlink() or not path.is_file():
            raise StructuredIngressError("SQLite source must be a regular non-symlink file")
        if not isinstance(source_ref, str) or not source_ref.strip():
            raise StructuredIngressError("source_ref must be a non-empty provenance reference")
        if not isinstance(table_name, str) or not table_name.strip():
            raise StructuredIngressError("table_name must be a non-empty allowlisted identifier")
        if not isinstance(columns, tuple) or not columns or any(
            not isinstance(column, str) or not column.strip() for column in columns
        ):
            raise StructuredIngressError("columns must be a non-empty tuple of identifiers")
        if len(set(columns)) != len(columns) or len(columns) > self._limits.maximum_columns:
            raise StructuredIngressError("columns contain duplicates or exceed the configured limit")
        main_database = _read_bounded_regular_file(path, self._limits.maximum_source_bytes)
        wal_path = Path(f"{path}-wal")
        if wal_path.exists() or wal_path.is_symlink():
            wal_content = _read_bounded_regular_file(
                wal_path,
                self._limits.maximum_source_bytes - len(main_database),
            )
        else:
            wal_content = b""
        if len(main_database) + len(wal_content) > self._limits.maximum_source_bytes:
            raise StructuredIngressError("SQLite source and WAL exceed the configured byte limit")
        uri = f"file:{path.resolve().as_posix()}?mode=ro"
        try:
            connection = sqlite3.connect(uri, uri=True)
            connection.execute("PRAGMA query_only=ON")
            table_row = connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name = ?",
                (table_name,),
            ).fetchone()
            if table_row is None:
                raise StructuredIngressError("requested table is not present in the SQLite source")
            available_columns = tuple(
                row[1] for row in connection.execute(f"PRAGMA table_info({_quote_identifier(table_name)})")
            )
            if any(column not in available_columns for column in columns):
                raise StructuredIngressError("requested columns are not present in the allowlisted table")
            selected = ", ".join(_quote_identifier(column) for column in columns)
            query = f"SELECT {selected} FROM {_quote_identifier(table_name)} LIMIT ?"
            cursor = connection.execute(query, (self._limits.maximum_rows + 1,))
            rows = tuple((row_number + 2, tuple(row)) for row_number, row in enumerate(cursor.fetchall()))
        except StructuredIngressError:
            raise
        except sqlite3.Error as error:
            raise StructuredIngressError("SQLite source could not be safely read") from error
        finally:
            if "connection" in locals():
                connection.close()
        if len(rows) > self._limits.maximum_rows:
            raise StructuredIngressError("SQLite table exceeds the configured row limit")
        canonical_rows = [
            [_json_value(value, maximum_cell_characters=self._limits.maximum_cell_characters) for value in row]
            for _, row in rows
        ]
        extraction_manifest = {
            "database_sha256": hashlib.sha256(main_database + b"\0" + wal_content).hexdigest(),
            "table_name": table_name,
            "columns": list(columns),
            "rows": canonical_rows,
        }
        extraction_bytes = json.dumps(
            extraction_manifest,
            allow_nan=False,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        digest = hashlib.sha256(extraction_bytes).hexdigest()
        records = _records_from_rows(
            headers=columns,
            rows=rows,
            source_ref=source_ref,
            source_sha256=digest,
            sheet_name=table_name,
            limits=self._limits,
        )
        return StructuredExtraction(StructuredFormat.SQLITE, source_ref, digest, records)

    def _extract_csv(self, content: bytes, source_ref: str, digest: str) -> tuple[ExtractedRecord, ...]:
        try:
            text = content.decode("utf-8", errors="strict")
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)
            header_values = next(reader)
            headers = _headers(header_values, maximum_columns=self._limits.maximum_columns)
            rows = []
            for values in reader:
                if not values or all(not value for value in values):
                    continue
                rows.append((reader.line_num, tuple(values)))
                if len(rows) > self._limits.maximum_rows:
                    raise StructuredIngressError("CSV source exceeds the configured row limit")
        except StructuredIngressError:
            raise
        except (UnicodeError, csv.Error, StopIteration) as error:
            raise StructuredIngressError("CSV source is malformed or has no header") from error
        return _records_from_rows(
            headers=headers,
            rows=rows,
            source_ref=source_ref,
            source_sha256=digest,
            sheet_name=None,
            limits=self._limits,
        )

    def _extract_jsonl(self, content: bytes, source_ref: str, digest: str) -> tuple[ExtractedRecord, ...]:
        try:
            text = content.decode("utf-8", errors="strict")
        except UnicodeDecodeError as error:
            raise StructuredIngressError("JSONL source is not valid UTF-8") from error
        headers: tuple[str, ...] | None = None
        rows: list[tuple[int, tuple[object, ...]]] = []
        for line_number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(
                    line,
                    object_pairs_hook=_strict_object,
                    parse_constant=lambda token: (_ for _ in ()).throw(
                        StructuredIngressError(f"invalid JSON numeric constant: {token}")
                    ),
                )
            except StructuredIngressError:
                raise
            except (json.JSONDecodeError, ValueError) as error:
                raise StructuredIngressError(f"JSONL line {line_number} is malformed") from error
            if not isinstance(record, dict) or not record:
                raise StructuredIngressError(f"JSONL line {line_number} must be a non-empty object")
            current_headers = _headers(record.keys(), maximum_columns=self._limits.maximum_columns)
            if headers is None:
                headers = current_headers
            if set(current_headers) != set(headers):
                raise StructuredIngressError("JSONL records must have one consistent field set")
            rows.append((line_number, tuple(record[name] for name in headers)))
            if len(rows) > self._limits.maximum_rows:
                raise StructuredIngressError("JSONL source exceeds the configured row limit")
        if headers is None:
            raise StructuredIngressError("JSONL source contains no records")
        return _records_from_rows(
            headers=headers,
            rows=rows,
            source_ref=source_ref,
            source_sha256=digest,
            sheet_name=None,
            limits=self._limits,
        )

    def _extract_xlsx(
        self,
        content: bytes,
        source_ref: str,
        digest: str,
        sheet_name: str | None,
    ) -> tuple[ExtractedRecord, ...]:
        if sheet_name is not None and (not isinstance(sheet_name, str) or not sheet_name.strip()):
            raise StructuredIngressError("sheet_name must be a non-empty string or None")
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            if sheet_name is None:
                if len(workbook.sheetnames) != 1:
                    raise StructuredIngressError("multi-sheet workbooks require an explicit worksheet name")
                sheet_name = workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise StructuredIngressError("requested worksheet is not present")
            worksheet = workbook[sheet_name]
            if worksheet.max_row is None or worksheet.max_column is None:
                raise StructuredIngressError("worksheet has no declared dimensions")
            if worksheet.max_row > self._limits.maximum_rows + 1:
                raise StructuredIngressError("worksheet exceeds the configured row limit")
            if worksheet.max_column > self._limits.maximum_columns:
                raise StructuredIngressError("worksheet exceeds the configured column limit")
            row_iterator = worksheet.iter_rows()
            header_cells = next(row_iterator)
            headers = _headers((cell.value for cell in header_cells), maximum_columns=self._limits.maximum_columns)
            rows = []
            for row_number, cells in enumerate(row_iterator, start=2):
                values = tuple(cell.value for cell in cells)
                if not any(value is not None for value in values):
                    continue
                rows.append((row_number, values[: len(headers)]))
                if len(rows) > self._limits.maximum_rows:
                    raise StructuredIngressError("worksheet exceeds the configured row limit")
        except StructuredIngressError:
            raise
        except Exception as error:
            raise StructuredIngressError("XLSX source is malformed or could not be read") from error
        finally:
            if workbook is not None:
                workbook.close()
        return _records_from_rows(
            headers=headers,
            rows=rows,
            source_ref=source_ref,
            source_sha256=digest,
            sheet_name=sheet_name,
            limits=self._limits,
        )


def _quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'